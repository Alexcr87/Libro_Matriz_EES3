"""
export_excel.py - Exportación de matrícula a Excel con formato oficial
E.E.S N°3 - "Malvinas Argentinas"
"""

import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font, Alignment, Border, Side, PatternFill, GradientFill
)
from openpyxl.utils import get_column_letter

import database as db


def _thin_border():
    thin = Side(style='thin')
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _header_fill():
    return PatternFill("solid", fgColor="1A3A5C")


def export_curso_to_sheet(ws, curso, alumnos):
    """Escribe una hoja de matrícula con formato para un curso."""
    anio = curso['anio']
    div = curso['division']

    # ── Título ──────────────────────────────────────────────────────────────
    ws.merge_cells('A1:U1')
    ws['A1'] = "REGISTRO DE MATRÍCULA"
    ws['A1'].font = Font(name='Calibri', bold=True, size=14, color="FFFFFF")
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws['A1'].fill = PatternFill("solid", fgColor="1A3A5C")

    ws.merge_cells('A2:U2')
    ws['A2'] = 'E.E.S N°3 - "MALVINAS ARGENTINAS"'
    ws['A2'].font = Font(name='Calibri', bold=True, size=12, color="FFFFFF")
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws['A2'].fill = PatternFill("solid", fgColor="1A3A5C")

    ws['A3'] = "AÑO:"
    ws['B3'] = anio
    ws['C3'] = "DIV:"
    ws['D3'] = div
    ws['O3'] = "AÑO LECTIVO:"
    ws['P3'] = datetime.now().year
    ws['J3'] = "Res. 302/12"
    for cell in ['A3', 'B3', 'C3', 'D3', 'O3', 'P3', 'J3']:
        ws[cell].font = Font(name='Calibri', bold=True, size=10)

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18

    # ── Cabeceras ────────────────────────────────────────────────────────────
    headers = [
        "N° Ord.", "Fecha Inscr.", "N° Matríc.", "Apellido y Nombre",
        "Fecha Nac.", "Edad", "Sexo", "Nacionalidad", "Tipo y N° Doc.",
        "Nombre Tutor/Encargado", "Nac. Tutor", "Profesión Tutor",
        "Domicilio", "Tel.", "Egresó", "Pase", "Salida >30/4",
        "Ingresa", "Ingr. >30/4", "Observaciones"
    ]
    row = 4
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(name='Calibri', bold=True, size=8, color="FFFFFF")
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.fill = PatternFill("solid", fgColor="2E6DA4")
        cell.border = _thin_border()
    ws.row_dimensions[row].height = 32

    # ── Datos ────────────────────────────────────────────────────────────────
    col_widths = [6, 12, 12, 28, 12, 5, 6, 12, 15, 28, 10, 14, 24, 12,
                  8, 8, 10, 10, 10, 22]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for r_idx, alumno in enumerate(alumnos, start=1):
        row_num = row + r_idx
        fill_color = "EBF3FB" if r_idx % 2 == 0 else "FFFFFF"
        fill = PatternFill("solid", fgColor=fill_color)

        values = [
            alumno['nro_orden'], alumno['fecha_inscripcion'],
            alumno['nro_matricula'], alumno['apellido_nombre'],
            alumno['fecha_nac'], alumno['edad'], alumno['sexo'],
            alumno['nacionalidad'], f"{alumno['tipo_doc'] or ''} {alumno['nro_doc'] or ''}".strip(),
            alumno['nombre_tutor'], alumno['nacionalidad_tutor'],
            alumno['profesion_tutor'], alumno['domicilio'],
            alumno['telefono'], alumno['egreso'], alumno['pase'],
            alumno['salida_despues_30_4'], alumno['ingresa'],
            alumno['ingr_despues_30_4'], alumno['observaciones']
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = Font(name='Calibri', size=8)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.fill = fill
            cell.border = _thin_border()
        ws.row_dimensions[row_num].height = 14

    # Filas vacías hasta 40
    for extra in range(len(alumnos) + 1, 41):
        row_num = row + extra
        fill_color = "EBF3FB" if extra % 2 == 0 else "FFFFFF"
        fill = PatternFill("solid", fgColor=fill_color)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.fill = fill
            cell.border = _thin_border()
            cell.font = Font(name='Calibri', size=8)
        ws.row_dimensions[row_num].height = 14


def export_all_cursos(filepath=None):
    """Exporta todos los cursos a un archivo Excel."""
    if not filepath:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filepath = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            f"Matricula_Export_{ts}.xlsx"
        )

    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    cursos = db.get_all_cursos()
    for curso in cursos:
        sheet_name = f"{curso['anio']}° {curso['division']}"
        ws = wb.create_sheet(title=sheet_name)
        alumnos = db.get_alumnos_by_curso(curso['id'])
        export_curso_to_sheet(ws, curso, alumnos)

    wb.save(filepath)
    return filepath


def export_single_curso(curso_id, filepath=None):
    """Exporta un único curso a Excel."""
    conn = db.get_connection()
    curso = conn.execute("SELECT * FROM cursos WHERE id=?", (curso_id,)).fetchone()
    conn.close()

    if not filepath:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        nombre = f"Matricula_{curso['anio']}{curso['division']}_{ts}.xlsx"
        filepath = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))), nombre
        )

    wb = Workbook()
    ws = wb.active
    ws.title = f"{curso['anio']}° {curso['division']}"
    alumnos = db.get_alumnos_by_curso(curso_id)
    export_curso_to_sheet(ws, curso, alumnos)
    wb.save(filepath)
    return filepath
